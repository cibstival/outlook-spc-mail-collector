# -*- coding: utf-8 -*-
"""SPC 폴더 넣기 규칙 - 사용자가 실제로 저장한 위치에서 배운다(폴더 이름 추측 대신).

방법
  수집함에 모인 메일 첨부(과거 메일 포함)와 원래 SPC 폴더의 파일을 '이름+크기' 또는 '내용(앞 1MB md5)'으로 짝짓는다.
  짝이 된 파일마다 사용자가 둔 위치(SPC 안 상대 경로)를 자료 기간 기준 토큰으로 바꿔 규칙 후보를 센다:
    AT/[AT]한빛_8기(2026.06)/회사제시  →  AT/[AT]한빛_{N}기({YYYY.MM+0})/회사제시
  (SPC, 종류)마다 가장 많이 쓴 틀이 3회 이상·80% 이상이면 '확실'.

넣을 때(수집기)
  - 설정 place_in_spc가 true이고, 그 SPC가 설정 place_allow 목록에 있고, (SPC, 종류) 규칙이 '확실'일 때만.
  - 규칙을 그 결산 기간으로 채운 폴더가 **이미 있을 때만** 넣는다(사용자가 그 기간 작업 폴더를 만든 뒤). 없으면 수집함에 두고,
    나중 실행에서 폴더가 생기면 옮긴다(promote).
  - 덮어쓰지 않음, 배치로그 기록.

결과 파일: 수집/spc_rules.json(로컬 전용 - SPC 이름이 들어가므로 저장소 제외)
"""
import os, re, sys, json, hashlib, datetime as dt
from collections import Counter, defaultdict

HERE = os.path.dirname(os.path.abspath(__file__))
SKIP = re.compile(r"과거|백업|(^|[\\/])old([\\/]|$)|사본|폐기|실수|자동화버전|표준화버전", re.I)
P_DATE = re.compile(r"(?<!\d)(20[12]\d)([.\-_ ]?)(0[1-9]|1[0-2])(?:\2([0-3]\d))?(?!\d)")
P_YY = re.compile(r"(?<![\d.])(\d\d)\.(0[1-9]|1[0-2])(?![\d.])")
P_K = re.compile(r"(?<!\d)(\d{1,2})(\s*기)")
TOKEN = re.compile(r"\{(YYYY(?P<s1>[.\-_ ]?)MM(?:(?P=s1)DD)?|YY\.MM)(?P<off>[+-]\d+)\}")
MIN_N, MIN_SHARE = 3, 0.8
NO_PLACE_KINDS = {"기타"}                     # 종류를 모르는 파일은 넣지 않음
MONTHLY_KINDS = {"자금판", "거래내역", "잔액증명", "발행내역"}   # 매달 오는 자료 - 결산 기간 폴더에 모으는 경우가 많아 결산월 기준으로 셈


def closing_period(y, m, months):
    """(년,월) → 결산월 목록 중 같거나 다음 결산월"""
    if not months:
        return y, m
    for k in range(13):
        yy, mm = _ym(y, m, k)
        if mm in months:
            return yy, mm
    return y, m


def _ym(y, m, off=0):
    t = y * 12 + (m - 1) + off
    return t // 12, t % 12 + 1


def seg_template(seg, base):
    """폴더 이름 한 칸 → 자료 기간(base=(년,월)) 대비 토큰"""
    m = P_DATE.search(seg)
    if m:
        y, mo = int(m.group(1)), int(m.group(3))
        off = (y * 12 + mo) - (base[0] * 12 + base[1])
        sep = m.group(2)
        fmt = "YYYY" + sep + "MM" + (sep + "DD" if m.group(4) else "")
        seg = seg[:m.start()] + "{" + fmt + f"{off:+d}" + "}" + seg[m.end():]
    else:
        m = P_YY.search(seg)
        if m:
            y, mo = 2000 + int(m.group(1)), int(m.group(2))
            off = (y * 12 + mo) - (base[0] * 12 + base[1])
            seg = seg[:m.start()] + "{YY.MM" + f"{off:+d}" + "}" + seg[m.end():]
    return P_K.sub(lambda k: "{N}" + k.group(2), seg)


def _md5(p, n=1 << 20):
    h = hashlib.md5()
    with open(p, "rb") as f:
        h.update(f.read(n))
    return h.hexdigest()


def _walk(root, skip_dir=None):
    for d, ds, fs in os.walk(root):
        if skip_dir:
            ds[:] = [x for x in ds if not skip_dir(x)]
        for f in fs:
            p = os.path.join(d, f)
            try:
                yield p, f, os.path.getsize(p)
            except OSError:
                continue


def learn(spc_root, inbox_root, kind_of, out=None, log=print, spc_files=None, scope=None):
    """수집함 ↔ SPC 폴더 짝짓기로 규칙 학습 → dict 저장. spc_files: 미리 만든 [(경로, 크기)](없으면 직접 훑음).
    scope(업무배정표): 있으면 배정표 회사만 배우고, 매달 오는 자료는 그 회사 결산월 기준으로 셈."""
    by_size = defaultdict(list)
    src = ((p, os.path.basename(p), s) for p, s in spc_files) if spc_files is not None else _walk(spc_root)
    for p, f, s in src:
        if s > 2000:
            by_size[s].append(p)
    rules = defaultdict(Counter)
    examples = defaultdict(list)
    n_pair = 0
    for p, f, s in _walk(inbox_root, skip_dir=lambda x: x.startswith("_")):
        if f.startswith("_"):
            continue
        rel_in = os.path.relpath(p, inbox_root).replace("\\", "/").split("/")
        per = next((x for x in reversed(rel_in[:-1]) if re.fullmatch(r"\d\d\.\d\d", x)), None)
        cands = by_size.get(s, [])
        if not per or not cands:
            continue
        same = [c for c in cands if os.path.basename(c) == f]
        if not same:
            ext = os.path.splitext(f)[1].lower()
            c2 = [c for c in cands if os.path.splitext(c)[1].lower() == ext][:6]
            if c2:
                try:
                    h = _md5(p)
                    same = [c for c in c2 if _md5(c) == h]
                except OSError:
                    same = []
        if not same:
            continue
        n_pair += 1
        base0 = (2000 + int(per[:2]), int(per[3:]))
        k = kind_of(f)
        for up in same:
            u = os.path.relpath(up, spc_root).replace("\\", "/").split("/")
            if len(u) < 3:
                continue
            spc_key, rel = "/".join(u[:2]), u[2:-1]
            ent = scope.match(u[1]) if scope else None
            if scope and not ent:
                continue
            base = closing_period(*base0, ent["결산월"]) if (ent and k in MONTHLY_KINDS) else base0
            if any(SKIP.search(x) for x in rel):
                continue
            tpl = "/".join(seg_template(x, base) for x in rel)
            rules[f"{spc_key}|{k}"][tpl] += 1
            if len(examples[f"{spc_key}|{k}|{tpl}"]) < 2:
                examples[f"{spc_key}|{k}|{tpl}"].append("/".join(rel))
    res = {}
    for key, c in rules.items():
        n = sum(c.values())
        top, tn = c.most_common(1)[0]
        sure = n >= MIN_N and tn / n >= MIN_SHARE and key.split("|")[1] not in NO_PLACE_KINDS and top != ""
        res[key] = {"틀": top, "횟수": tn, "전체": n, "확실": bool(sure), "다른틀": dict(c.most_common()[1:4]),
                    "예": examples.get(f"{key}|{top}", [])}
    data = {"만든때": f"{dt.datetime.now():%Y-%m-%d %H:%M}", "짝": n_pair, "규칙": res}
    if out:
        json.dump(data, open(out, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    log(f"짝지은 파일 {n_pair}개, 규칙 {len(res)}개(확실 {sum(r['확실'] for r in res.values())}개)")
    return data


def _seg_regex(pat):
    """틀 한 칸 → 정규식({N}→기수, 날짜 토큰→날짜 글자)"""
    out, i = [], 0
    for m in re.finditer(r"\{N\}|" + TOKEN.pattern, pat):
        out.append(re.escape(pat[i:m.start()]))
        out.append(r"(?P<N>\d{1,2})" if m.group(0) == "{N}" else r"(?P<D>\d{2,4}[.\-_ ]?\d{2}(?:[.\-_ ]?\d{2})?)")
        i = m.end()
    out.append(re.escape(pat[i:]))
    return re.compile("^" + "".join(out) + "$")


def _kisu_map(spc_dir, tpl_parts, _unused=None):
    """틀과 같은 모양의 기존 폴더에서 (년,월) → 기수 표본"""
    out = {}
    def rec(d, i):
        if i >= len(tpl_parts) or not os.path.isdir(d):
            return
        rx = _seg_regex(tpl_parts[i])
        for name in os.listdir(d):
            full = os.path.join(d, name)
            if not os.path.isdir(full):
                continue
            m = rx.match(name)
            if not m:
                continue
            g = m.groupdict()
            if g.get("N") and g.get("D"):
                dm = P_DATE.search(g["D"])
                ym = None if dm else P_YY.search(g["D"])
                if dm:
                    out[(int(dm.group(1)), int(dm.group(3)))] = int(g["N"])
                elif ym:
                    out[(2000 + int(ym.group(1)), int(ym.group(2)))] = int(g["N"])
            rec(full, i + 1)
    rec(spc_dir, 0)
    return out


def _kisu(samples, y, m):
    if not samples:
        return None
    if (y, m) in samples:
        return samples[(y, m)]
    ks = sorted(samples.items())
    ends = set()                                   # 기수가 바뀌는 직전 달 = 사업연도 끝
    for (a, ka), (b, kb) in zip(ks, ks[1:]):
        if kb == ka + 1:
            ends.add(a[1])
    ends = ends or {12}
    (y0, m0), k0 = min(ks, key=lambda x: abs((x[0][0] * 12 + x[0][1]) - (y * 12 + m)))
    k, cy, cm = k0, y0, m0
    step = 1 if (y, m) > (y0, m0) else -1
    while (cy, cm) != (y, m):
        if step > 0 and cm in ends:
            k += 1
        cy, cm = _ym(cy, cm, step)
        if step < 0 and cm in ends:
            k -= 1
    return k


def target_dir(spc_dir, tpl, per):
    """규칙 틀 + 결산 기간('YY.MM') → SPC 안 폴더 경로(이미 있을 때만), 없으면 None"""
    y, m = 2000 + int(per[:2]), int(per[3:])
    parts = tpl.split("/")
    def fill(seg, kisu):
        def tok(mm):
            off = int(mm.group("off"))
            yy, mo = _ym(y, m, off)
            f = mm.group(1)
            if f == "YY.MM":
                return f"{yy % 100:02d}.{mo:02d}"
            sep = mm.group("s1") or ""
            last = (dt.date(yy + mo // 12, mo % 12 + 1, 1) - dt.timedelta(days=1)).day
            return f"{yy}{sep}{mo:02d}" + (f"{sep}{last:02d}" if "DD" in f else "")
        seg = TOKEN.sub(tok, seg)
        if "{N}" in seg:
            if kisu is None:
                return None
            seg = seg.replace("{N}", str(kisu))
        return seg
    kisu = None
    if any("{N}" in p for p in parts):
        kisu = _kisu(_kisu_map(spc_dir, parts, None), y, m)
    path = spc_dir
    for p in parts:
        seg = fill(p, kisu)
        if seg is None:
            return None
        path = os.path.join(path, seg)
    return path if os.path.isdir(path) else None


def preview(rules, spc_root, out_xlsx, periods=("26.06", "26.09", "26.12"), scope=None):
    """SPC·종류별 규칙과 근거, 다음 기간 위치(폴더 있으면) - 사용자 확인용 엑셀(로컬 전용)"""
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "넣을 위치(확인용)"
    ws.append(["메일로 받은 자료를 예전에 직접 저장하신 위치에서 배운 규칙입니다. '확실'만 쓰고, 그 기간 폴더가 이미 있을 때만 넣습니다."])
    ws.append(["허용할 SPC는 '허용' 칸에 O 를 적어 Claude에게 알려주세요(또는 그대로 두면 아무 데도 넣지 않음)."])
    head = ["허용", "수탁사/SPC", "종류", "확실", "근거(같은 곳/전체)", "배운 틀", "예전 실제 위치(예)"] + [f"{p} 위치" for p in periods]
    ws.append(head)
    for c in ws[3]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    for key, r in sorted(rules["규칙"].items(), key=lambda x: (x[0].split("|")[0], -x[1]["전체"])):
        spc_key, kind = key.split("|")
        if scope and not scope.match(spc_key.split("/")[-1]):
            continue
        cells = []
        for per in periods:
            if not r["확실"]:
                cells.append("")
                continue
            t = target_dir(os.path.join(spc_root, *spc_key.split("/")), r["틀"], per)
            cells.append(os.path.relpath(t, os.path.join(spc_root, *spc_key.split("/"))) if t else "(폴더 아직 없음 → 수집함)")
        ws.append(["", spc_key, kind, "확실" if r["확실"] else ("애매" if r["전체"] >= MIN_N else "표본 적음"),
                   f"{r['횟수']}/{r['전체']}", r["틀"], " | ".join(r["예"])] + cells)
        if not r["확실"]:
            for c in ws[ws.max_row]:
                c.font = Font(color="888888")
    for col, w in zip("ABCDEFGHIJ", (6, 40, 10, 8, 12, 50, 50, 30, 30, 30)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"
    wb.save(out_xlsx)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.path.insert(0, HERE)
    import ingest_spc as I
    a = sys.argv[1:]
    cfg = json.load(open(a[a.index("--config") + 1] if "--config" in a else os.path.join(HERE, "spc_mail_config.json"), encoding="utf-8"))
    out = os.path.join(HERE, "spc_rules.json")
    import scope as S
    sc = S.load(cfg)
    data = learn(cfg["spc_root"], cfg["dest_root"], lambda f: I.Router.kind(f, "", "")[0], out, scope=sc)
    xl = os.path.join(os.path.dirname(HERE), "SPC폴더_넣을위치_확인.xlsx")
    preview(data, cfg["spc_root"], xl, scope=sc)
    print("확인용 엑셀:", xl)
